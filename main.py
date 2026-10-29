import os
from pathlib import Path 
from dotenv import load_dotenv
from pydantic import BaseModel
from groq import Groq
from pypdf import PdfReader
import time
from docx import Document
import json


load_dotenv()
my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("API key not found")

client= Groq(api_key=my_api_key)

model = "openai/gpt-oss-120b"
role="user"

class JobDescription(BaseModel):
    Role:str
    Required_skills:list[str]
    Prefferd_skills:list[str]
    Min_experience:int | None
    Education:str
    responsibility:str
class Experience(BaseModel):
    Company:str | None=None
    Role:str | None=None
    Duration:str | None=None
    Description:str | None=None
    Skills:list[str] = []

class Resume(BaseModel):
    Name:str | None=None
    Email:str | None=None
    Phone:int | None=None
    Total_experience_Year:float | None=None
    Skills:list[str] = []
    Experiences:list[Experience] = []
    Projects:list[str] = []
    Certifications:list[str] = []

class MatchResult(BaseModel):
    score:float
    details:dict

JobDescription_schema = JobDescription.model_json_schema()
Experience_schema =Experience.model_json_schema()
Resume_schema = Resume.model_json_schema()

Job_Description_str='''Amazon is looking for a Software Development Engineer I (SDE I) responsible for designing, developing, testing, 
debugging, deploying, and maintaining scalable software applications and services. Required skills include proficiency in
at least one programming language such as Java, Python, or C++, strong knowledge of data structures and algorithms, object-oriented programming, 
problem-solving, and computer science fundamentals. Preferred skills include experience with SQL or NoSQL databases, REST APIs, AWS cloud services,
Git, automated testing, and distributed systems. Educational requirements typically include a Bachelor's degree in Computer Science, Information Technology,
Computer Engineering, or a related field, or equivalent qualifications where accepted. Minimum experience varies by job posting; entry-level and 
university-hiring roles may accept fresh graduates, while some SDE I openings require one or more years of professional software development experience. Candidates should demonstrate
analytical thinking, communication, teamwork, ownership, attention to code quality, and a willingness to learn
and follow Amazon's Leadership Principles.'''
system_prompt=f'''
You are an expert HR assistant.
Your job is to analyze job descriptions and extract
structured information from them.
Return ONLY valid JSON matching this schema:
{JobDescription_schema}
IMPORTANT:
Do NOT return the schema itself.
Do NOT return fields like "properties", "title" or "type".
Fill the schema with actual information extracted from the job description.

If minimum experience is not mentioned, return null.
If information for a list is missing, return an empty list.
Do not invent information.
'''

user_prompt = f"""
Analyze the following job description:
{Job_Description_str}
"""

message_system = {
    "role": "system",
    "content": system_prompt
}

message_user = {
    "role": "user",
    "content": user_prompt
}

response_format = {
    "type": "json_object"
}

messages = [message_system, message_user]
response = client.chat.completions.create(model=model,messages=messages,response_format=response_format)
answer = response.choices[0].message.content
raw_json = answer
jobdata = json.loads(raw_json)
Job_D = JobDescription(**jobdata)

#resume parser

def parse_resume(resume_text):
    system_prompt=f"""
You are an expert resume parser.
Extract information from the resume based on its meaning,
not only based on exact section headings.
Different resumes may use different headings.
For example:
- Experience
- Professional Experience
- Work History
- Employment
- Internships
These may all contain relevant experience.
Skills may also appear in the skills section, work experience,
internships or projects.
Return ONLY valid JSON matching this schema:
{Resume_schema}
Important rules:
1. Do not invent information.
2. If a value is not available, return null.
3. If a list has no information, return an empty list.
4. Include internships inside experiences.
5. Extract skills mentioned across the entire resume.
"""
    user_prompt=f"""parse the following resume {resume_text}"""
    message_system = {
        "role":"system",
        "content":system_prompt
    }
    message_user = {
        "role":"user",
        "content":user_prompt
    }
    messages = [message_system,message_user]
    response_format = {
        "type":"json_object"
    }

    raw = client.chat.completions.create(
        model=model,
        messages = messages,
        response_format=response_format
    ).choices[0].message.content
    data = json.loads(raw)
    resume = Resume(**data)
    return resume

def final_score(job, resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
You are an HR recruiter.

Compare the candidate's resume with the job description.

JOB DESCRIPTION:
{job.model_dump_json(indent=2)}

CANDIDATE RESUME:
{resume.model_dump_json(indent=2)}
Return JSON matching this schema:

{match_schema}

Give me:

1. Candidate name
2. Matching skills
3. Missing important skills
4. Whether experience requirement is met
5. Overall match percentage from 0 to 100
6. A short final verdict
Keep the response concise and easy to read.
"""

    message = {
        "role": "user",
        "content": prompt
    }
    messages = [message]

    response_format = {
        "type": "json_object"
    }

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format=response_format
    )

    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)

def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    pgtext=""
    for page in reader.pages:
        pgtext += page.extract_text()
        if pgtext:
            text += pgtext + '\n'
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for para in document.paragraphs:
        if para.text.strip():
            text += para.text + "\n"

    for table in document.tables:
        for row in table.row:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None


resume_folder = Path("resume")
all_result = []
for file_path in resume_folder.iterdir():
    # file path contain all paths of files
    if file_path.suffix.lower() not in [".pdf",".docx"]:
        continue 
    print("\nProcessing: ",file_path)
    Resume_text = read_resume(file_path)
    parsed_resume = parse_resume(Resume_text);
    time.sleep(5) # rate limiting hence
    result = final_score(Job_D,parsed_resume)
    time.sleep(5)
    print("Score",result.score)
    all_result.append({
        "name":parsed_resume.Name,
        "score":result.score,
        "details":result.details
    })

all_result.sort(
    key=lambda candidate:candidate["score"],
    reverse=True
)
top_2 = all_result[:2]
wrost_2 = all_result[-2:]

print("Top 2")
for candidate in top_2:
    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )
    print(candidate["details"])

print("Lowest 2")
for candidate in wrost_2:
    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )
    print(candidate["details"])
    